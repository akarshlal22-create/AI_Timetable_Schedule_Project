"""tests/test_exporter.py — Tests for CSV and Excel export functionality."""

from __future__ import annotations

import io
import os
import sys

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

import openpyxl
import pandas as pd
import pytest

from algorithms.backtracking import BacktrackingConfig, BacktrackingSolver
from core.loader import load_problem_from_json
from core.models import Timetable
from services.exporter import (
    build_flat_sessions_dataframe,
    export_flat_csv,
    export_grid_csv,
    export_timetable_excel,
)
from services.grid_builder import build_grid_dataframe


@pytest.fixture
def solved_timetable() -> Timetable:
    data_path = os.path.join(os.path.dirname(__file__), "..", "data", "small.json")
    p = load_problem_from_json(data_path)
    solver = BacktrackingSolver(
        BacktrackingConfig(use_mrv=True, use_degree=True, use_forward_checking=True)
    )
    result = solver.solve(p, seed=42)
    assert result.timetable is not None
    return result.timetable


def test_export_grid_csv(solved_timetable: Timetable):
    grid_df, _ = build_grid_dataframe(solved_timetable, view_type="class", entity_name="CS-A")
    csv_str = export_grid_csv(grid_df)

    assert isinstance(csv_str, str)
    lines = csv_str.strip().splitlines()
    # 1 header line + 6 period rows
    assert len(lines) == 7


def test_build_flat_sessions_dataframe(solved_timetable: Timetable):
    df = build_flat_sessions_dataframe(solved_timetable)

    expected_cols = ["Class", "Subject", "Teacher", "Room", "Day", "Period", "Is_Lab"]
    assert list(df.columns) == expected_cols
    assert len(df) == len(solved_timetable.assignment)


def test_export_flat_csv(solved_timetable: Timetable):
    csv_str = export_flat_csv(solved_timetable)
    assert isinstance(csv_str, str)

    df = pd.read_csv(io.StringIO(csv_str))
    expected_cols = ["Class", "Subject", "Teacher", "Room", "Day", "Period", "Is_Lab"]
    assert list(df.columns) == expected_cols
    assert len(df) == len(solved_timetable.assignment)


def test_export_timetable_excel(solved_timetable: Timetable):
    excel_bytes = export_timetable_excel(solved_timetable)
    assert isinstance(excel_bytes, bytes)
    assert len(excel_bytes) > 0

    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    sheet_names = wb.sheetnames

    p = solved_timetable.problem

    # Verify Class sheets
    for cls_name in p.classes.keys():
        expected_title = f"Class_{cls_name}"[:31]
        assert expected_title in sheet_names

    # Verify Teacher sheets
    for t_name in p.teachers.keys():
        clean_name = t_name.replace(":", "_").replace("/", "_").replace("\\", "_")
        expected_title = f"T_{clean_name}"[:31]
        assert expected_title in sheet_names

    # Verify Room sheets
    for r_name in p.rooms.keys():
        expected_title = f"Room_{r_name}"[:31]
        assert expected_title in sheet_names

    # Verify All_Sessions sheet
    assert "All_Sessions" in sheet_names

    # Verify All_Sessions has data rows
    ws_all = wb["All_Sessions"]
    assert ws_all.max_row == len(solved_timetable.assignment) + 1  # header + rows

